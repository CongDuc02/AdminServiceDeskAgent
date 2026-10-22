"""Dấu giữ chỗ của biến điền-sau-duyệt và đọc chữ của `.docx` (A-098; PO 2026-10-15) — hàm thuần, không hạ tầng.

Phép kiểm "bản FINAL không còn dấu giữ chỗ nào" ở mức tool (PostgreSQL, SeaweedFS, `soffice`) là `tests/test_render_tools.py`, lớp `FinalKhongConDauGiuCho`.
Chạy từ backend/:  PYTHONPATH=src python -m unittest tests.test_provisional -v
"""
from __future__ import annotations

from tests import _guard  # noqa: F401 — chốt chặn mạng và khoá API của bộ test (tests/_guard.py)

import io
import unicodedata
import unittest

from docx import Document

from bo19.domain import provisional
from bo19.orchestrator.document_graph import failure
from bo19.tool_layer.rendering import doc_text
from tests.render_support import build_template


class Marker(unittest.TestCase):
    def test_dang_nhin_ra_ngay_la_chua_chinh_thuc(self):
        self.assertEqual(provisional.marker_for("document_number"), "[Số văn bản — cấp khi phát hành]")
        self.assertEqual(provisional.marker_for("issued_date"), "[Ngày ban hành — cấp khi phát hành]")
        self.assertEqual(provisional.marker_for("signer_user_id"), "[Người ký — cấp khi phát hành]")
        self.assertEqual(provisional.marker_for("bien_la"), "[bien_la — cấp khi phát hành]")  # biến lạ: hiện tên biến, vẫn đúng khuôn

    def test_moi_dau_do_he_thong_sinh_ra_deu_nhan_ra_duoc(self):
        for name in ("document_number", "issued_date", "signer_user_id", "bien_la", "a"):
            self.assertTrue(provisional.contains_marker(provisional.marker_for(name)), name)

    def test_nhan_ra_theo_khuon_khong_phan_biet_hoa_thuong_gach_va_dang_unicode(self):
        for text in ("[số văn bản - CẤP KHI PHÁT HÀNH]", "[Số văn bản – cấp khi phát hành]", "Số: [Số văn bản — cấp khi phát hành], ngày", "[Ngày — cấp khi phát hành ]",
                     unicodedata.normalize("NFD", provisional.marker_for("document_number"))):
            self.assertTrue(provisional.contains_marker(text), repr(text))

    def test_khong_nham_van_ban_that(self):
        for text in ("", "Số văn bản được cấp khi phát hành", "[Số văn bản]", "cấp khi phát hành]", "[— cấp khi phát hành]", "123/QĐ-GIẢ", "Điều [1] — cấp khi phát hành quyết định"):
            self.assertFalse(provisional.contains_marker(text), repr(text))

    def test_khong_nuot_qua_nhieu_dau_ngoac(self):
        self.assertFalse(provisional.contains_marker("[a] x — cấp khi phát hành]"))  # `[` mở trước `]` đầu: không phải một dấu

    def test_ten_bien_co_gia_tri_giu_cho_chi_tra_ten(self):
        values = {"a": "Nguyễn Thị Giang", "b": provisional.marker_for("document_number"), "c": "...", "d": "N/A", "e": "[điền sau]", "f": "  ", "g": "Trần Văn Bình"}
        self.assertEqual(provisional.placeholder_variables(values), ("b", "c", "d", "e", "f"))
        self.assertEqual(provisional.placeholder_variables({"a": "Nguyễn Thị Giang", "h": ""}), ())  # "" đúng nghĩa: slot không bắt buộc vắng (có chủ đích); "  " thì vẫn là dấu giữ chỗ ở trên
        self.assertEqual(provisional.placeholder_variables({}), ())


class DocText(unittest.TestCase):
    def test_noi_cac_run_cua_cung_mot_doan(self):
        document = Document()
        p = document.add_paragraph()
        p.add_run("Số: [Số văn bản — ")
        p.add_run("cấp khi ")
        p.add_run("phát hành]")
        document.add_paragraph("Đoạn hai")
        buf = io.BytesIO()
        document.save(buf)
        texts = doc_text.paragraph_texts(buf.getvalue())
        self.assertIn("Số: [Số văn bản — cấp khi phát hành]", texts)
        self.assertIn("Đoạn hai", texts)
        self.assertTrue(any(provisional.contains_marker(t) for t in texts))

    def test_doc_than_header_va_footer(self):
        data = build_template(header_text="ĐẦU TRANG GIẢ", footer_text="CHÂN TRANG GIẢ", extra_paragraphs=("Đoạn thêm & <ký hiệu>",))
        texts = doc_text.paragraph_texts(data)
        for expected in ("ĐẦU TRANG GIẢ", "CHÂN TRANG GIẢ", "{{ tieu_de }}", "Đoạn thêm & <ký hiệu>"):
            self.assertIn(expected, texts)  # `&`, `<` được trả về dạng chữ, không dạng thực thể XML

    def test_dau_trong_header_va_footer_duoc_thay(self):
        mark = provisional.marker_for("document_number")
        for kw in ({"header_text": mark}, {"footer_text": mark}, {"extra_paragraphs": (mark,)}):
            self.assertTrue(any(provisional.contains_marker(t) for t in doc_text.paragraph_texts(build_template(**kw))), kw)
        self.assertFalse(any(provisional.contains_marker(t) for t in doc_text.paragraph_texts(build_template())))


class MaLoi(unittest.TestCase):
    def test_con_dau_giu_cho_o_ban_final_la_loi_vinh_vien(self):
        """Thử lại cho cùng kết quả: `finalize_issue` (B9b) không được lặp vô ích."""
        self.assertEqual(failure.classify("PLACEHOLDER_REMAINS"), failure.Failure("PLACEHOLDER_REMAINS", True))
        self.assertEqual(failure.classify("PLACEHOLDER_REMAINS:document_text"), failure.Failure("PLACEHOLDER_REMAINS", True))

    def test_ma_cua_tool_khop_ma_trong_bang(self):
        from bo19.tool_layer.tools import docx_render
        self.assertEqual(docx_render.PlaceholderRemains.code, "PLACEHOLDER_REMAINS")


if __name__ == "__main__":
    unittest.main()

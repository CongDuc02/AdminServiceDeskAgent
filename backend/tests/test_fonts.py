"""`bo19.tool_layer.rendering.fonts` (B8b; ADR-015 chốt 1–3, ADR-036) và các file font của repo.

Điều được chứng minh, nói bằng hành vi:
- **họ font mà `.docx` khai** — đọc được từ từng nơi đã biết: đoạn chạy ở thân văn bản và ở tiêu đề, kiểu định dạng và mặc định, bảng font, ký hiệu `w:sym`, phông của theme, phần chú thích cuối trang; file hỏng là `DOCX_UNREADABLE`;
- **"có mặt"** — họ có ánh xạ tường minh thì `fc-match` phải ra ĐÚNG họ đích (không đủ là "có một font nào đó"); họ khác thì khớp chính xác tên, không khớp chuỗi con; công cụ thiếu là `FONT_TOOL_MISSING`;
- **file font ghim:** bốn file Liberation Serif khớp sha256 ở `docs/reference/font-liberation-serif-license.md`; có `LICENSE`; `fonts/50-bo19-times-new-roman.conf` khớp `FONT_ALIASES`, và QUY TẮC TRONG FILE ấy là thứ làm "Times New Roman" ra Liberation Serif (thử với cấu hình fontconfig cô lập, không dựa vào quy tắc thay thế có sẵn của hệ điều hành);
- **`Dockerfile`:** `runtime` là stage CUỐI (PO, 2026-10-13) và `test` không đứng sau nó.

Các ca cần `fc-match`/`fc-list` chạy trong image target `test`; thiếu thì bỏ qua — trừ khi `BO19_REQUIRE_SOFFICE=1` (CI).
Chạy từ backend/ với PYTHONPATH=src:  python -m unittest tests.test_fonts -v
"""
from __future__ import annotations

from tests import _guard  # noqa: F401 — chốt chặn mạng và khoá API của bộ test (tests/_guard.py)

import hashlib
import io
import os
import re
import subprocess
import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document
from lxml import etree

from bo19.tool_layer.rendering import fonts
from tests.render_support import build_template, require_tool

REPO = Path(__file__).resolve().parents[2]
FONT_DIR = REPO / "fonts"
REFERENCE = REPO / "docs" / "reference" / "font-liberation-serif-license.md"


def with_part(docx: bytes, name: str, xml: str) -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(docx)) as src, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            dst.writestr(item, src.read(item.filename))
        dst.writestr(name, xml)
    return out.getvalue()


W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'


class HoDaKhai(unittest.TestCase):
    def test_mau_khai_times_new_roman_o_kieu_mac_dinh_tieu_de_chan_trang(self):
        self.assertIn("Times New Roman", fonts.declared_families(build_template()))

    def test_doan_chay_o_than_van_ban(self):
        d = Document(io.BytesIO(build_template()))
        d.add_paragraph().add_run("x").font.name = "Garamond"
        buf = io.BytesIO()
        d.save(buf)
        self.assertIn("Garamond", fonts.declared_families(buf.getvalue()))

    def test_doan_chay_o_tieu_de_va_chan_trang(self):
        d = Document(io.BytesIO(build_template()))
        d.sections[0].header.paragraphs[0].add_run("h").font.name = "Arial Narrow"
        d.sections[0].footer.paragraphs[0].add_run("f").font.name = "Palatino"
        buf = io.BytesIO()
        d.save(buf)
        got = fonts.declared_families(buf.getvalue())
        self.assertLessEqual({"Arial Narrow", "Palatino"}, got)

    def test_kieu_dinh_dang_va_mac_dinh(self):
        d = Document(io.BytesIO(build_template()))
        d.styles.add_style("Kieu Rieng", 1).font.name = "Verdana"
        buf = io.BytesIO()
        d.save(buf)
        self.assertIn("Verdana", fonts.declared_families(buf.getvalue()))
        styles = zipfile.ZipFile(io.BytesIO(build_template())).read("word/styles.xml").decode("utf-8")
        self.assertIn("<w:docDefaults>", styles)  # docDefaults là một phần của styles.xml, cùng đường đọc

    def test_bang_font_va_ky_hieu(self):
        got = fonts.declared_families(with_part(build_template(), "word/extra.xml", f'<w:root {W}><w:font w:name="Wingdings 2"/><w:sym w:font="Webdings" w:char="F0A7"/></w:root>'))
        self.assertLessEqual({"Wingdings 2", "Webdings"}, got)

    def test_phong_cua_theme_khi_dan_bang_theme(self):
        theme = zipfile.ZipFile(io.BytesIO(build_template())).read("word/theme/theme1.xml").decode("utf-8")
        minor = re.search(r"<a:minorFont>\s*<a:latin typeface=\"([^\"]+)\"", theme).group(1)
        major = re.search(r"<a:majorFont>\s*<a:latin typeface=\"([^\"]+)\"", theme).group(1)
        d = Document(io.BytesIO(build_template()))
        got = fonts.declared_families(d_bytes(d))
        # python-docx mặc định dùng phông theme ở docDefaults/Heading: phải thấy phông thật, không phải tên "minorHAnsi"
        self.assertTrue({minor, major} & got, (minor, major, got))
        self.assertNotIn("minorHAnsi", got)

    def test_theme_dan_bang_attribute_theme_duoc_giai(self):
        xml = (f'<w:root {W}><w:rPr><w:rFonts w:asciiTheme="majorHAnsi" w:hAnsiTheme="majorHAnsi"/></w:rPr></w:root>')
        theme = ('<a:theme xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:themeElements><a:fontScheme name="x">'
                 '<a:majorFont><a:latin typeface="Cambria Math Giả"/></a:majorFont><a:minorFont><a:latin typeface="Corbel Giả"/></a:minorFont></a:fontScheme></a:themeElements></a:theme>')
        out = io.BytesIO()
        with zipfile.ZipFile(out, "w") as z:  # chỉ hai phần: theme và một phần dẫn bằng theme — không phần nào khác dẫn tới phông `minor`
            z.writestr("word/theme/theme1.xml", theme)
            z.writestr("word/probe.xml", xml)
        got = fonts.declared_families(out.getvalue())
        self.assertIn("Cambria Math Giả", got)
        self.assertNotIn("Corbel Giả", got)  # `minor` không được dẫn tới

    def test_phan_chu_thich_cuoi_trang(self):
        got = fonts.declared_families(with_part(build_template(), "word/footnotes.xml", f'<w:footnotes {W}><w:footnote><w:p><w:r><w:rPr><w:rFonts w:ascii="Century Giả"/></w:rPr></w:r></w:p></w:footnote></w:footnotes>'))
        self.assertIn("Century Giả", got)

    def test_thuoc_tinh_eastasia_va_cs(self):
        got = fonts.declared_families(with_part(build_template(), "word/p.xml", f'<w:r {W}><w:rFonts w:eastAsia="MS Mincho Giả" w:cs="Tahoma Giả"/></w:r>'))
        self.assertLessEqual({"MS Mincho Giả", "Tahoma Giả"}, got)

    def test_file_hong_la_mot_ma(self):
        for blob in (b"", b"khong phai zip", with_part(build_template(), "word/hong.xml", "<w:root")):
            with self.assertRaises(fonts.DocxUnreadable):
                fonts.declared_families(blob)

    def test_khong_doc_thuc_the_ngoai_va_chan_phan_qua_lon(self):
        xxe = '<?xml version="1.0"?><!DOCTYPE r [<!ENTITY e SYSTEM "file:///etc/hostname">]>' + f'<w:r {W}><w:rFonts w:ascii="&e;"/></w:r>'
        try:
            got = fonts.declared_families(with_part(build_template(), "word/xxe.xml", xxe))
        except fonts.DocxUnreadable:
            got = frozenset()  # thực thể không được phân giải: lxml từ chối — cũng là một cách không đọc file ngoài
        if os.path.exists("/etc/hostname"):
            with open("/etc/hostname") as f:
                self.assertNotIn(f.read().strip(), got)
        old = fonts.MAX_PART_BYTES
        fonts.MAX_PART_BYTES = 10
        try:
            with self.assertRaises(fonts.DocxUnreadable):
                fonts.declared_families(build_template())
        finally:
            fonts.MAX_PART_BYTES = old


def d_bytes(document) -> bytes:
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


class CoMat(unittest.TestCase):
    """Luật "có mặt" trên đầu ra GIẢ của fc-match/fc-list — không cần công cụ thật."""

    @staticmethod
    def runner(match: dict[str, str] | None = None, installed: str = "Liberation Serif\nDejaVu Sans,DejaVu Sans Light\nOpenSymbol\n"):
        def run(cmd: list[str]) -> str:
            if cmd[0] == "fc-match":
                return (match or {}).get(cmd[-1], "DejaVu Sans")
            return installed
        return run

    def test_ho_co_anh_xa_phai_ra_dung_ho_dich(self):
        self.assertTrue(fonts.is_present("Times New Roman", run=self.runner({"Times New Roman": "Liberation Serif"})))
        self.assertFalse(fonts.is_present("Times New Roman", run=self.runner({"Times New Roman": "DejaVu Serif"})))

    def test_ho_co_anh_xa_khong_duoc_coi_la_co_mat_chi_vi_co_font_cung_ten_trong_danh_sach(self):
        run = self.runner({"Times New Roman": "DejaVu Serif"}, installed="Times New Roman\nLiberation Serif\n")
        self.assertFalse(fonts.is_present("Times New Roman", run=run))  # có một font tên "Times New Roman" nhưng fc-match không ra Liberation Serif: ánh xạ hỏng là lỗi

    def test_ho_khac_khop_chinh_xac_ten(self):
        run = self.runner()
        self.assertTrue(fonts.is_present("Liberation Serif", run=run))
        self.assertTrue(fonts.is_present("DejaVu Sans Light", run=run))  # một font mang nhiều tên: mỗi tên là một họ
        self.assertFalse(fonts.is_present("DejaVu", run=run))  # không khớp chuỗi con
        self.assertFalse(fonts.is_present("liberation serif", run=run))  # phân biệt hoa thường: "khớp chính xác"
        self.assertFalse(fonts.is_present("Arial", run=run))  # fc-match có thể ra font gần nhất — không được tính

    def test_missing_theo_thu_tu_chu_cai_va_goi_fc_list_mot_lan(self):
        calls = []

        def run(cmd):
            calls.append(cmd[0])
            return self.runner({"Times New Roman": "Liberation Serif"})(cmd)
        self.assertEqual(fonts.missing(["Zapf", "Times New Roman", "Arial", "Liberation Serif"], run=run), ["Arial", "Zapf"])
        self.assertEqual(calls.count("fc-list"), 1)

    def test_cong_cu_thieu_la_loi_chu_khong_phai_font_thieu(self):
        def run(cmd):
            raise fonts.FontToolMissing
        with self.assertRaises(fonts.FontToolMissing):
            fonts.missing(["Times New Roman"], run=run)

    def test_lenh_that_bai_hay_khong_ton_tai_la_font_tool_missing(self):
        with self.assertRaises(fonts.FontToolMissing):
            fonts._run(["lenh-khong-ton-tai-bo19"])
        with self.assertRaises(fonts.FontToolMissing):
            fonts._run(["false"])

    def test_phan_tich_dau_ra(self):
        self.assertEqual(fonts._families("A\\-B,C D\n\nE\n"), {"A-B", "C D", "E"})


class FcThat(unittest.TestCase):
    """`fc-match`/`fc-list` thật của image (target `test`/`runtime`)."""

    @classmethod
    def setUpClass(cls):
        require_tool(cls, "fc-match")

    def test_times_new_roman_ra_liberation_serif(self):
        self.assertEqual(fonts.match_families("Times New Roman"), {"Liberation Serif"})

    def test_liberation_serif_nam_trong_danh_sach_cai(self):
        self.assertIn("Liberation Serif", fonts.installed_families())
        self.assertEqual(fonts.missing(["Times New Roman", "Liberation Serif"]), [])

    def test_font_khong_cai_la_thieu_du_fc_match_van_tra_ve_font_gan_nhat(self):
        self.assertTrue(fonts.match_families("Font Khong Ton Tai Bo19"))  # fc-match LUÔN trả về một cái gì đó — đó chính là chỗ hỏng im lặng mà ADR-015 chặn
        self.assertEqual(fonts.missing(["Font Khong Ton Tai Bo19"]), ["Font Khong Ton Tai Bo19"])

    def test_quy_tac_trong_file_conf_la_thu_lam_times_new_roman_ra_liberation_serif(self):
        """Cấu hình fontconfig CÔ LẬP: chỉ hai thư mục font (Liberation Serif và DejaVu) — không `30-metric-aliases.conf` của hệ điều hành, vốn cũng ánh xạ cùng tên. Có file của dự án thì ra Liberation Serif;
        bỏ quy tắc (đổi họ đích) thì không."""
        conf = (FONT_DIR / "50-bo19-times-new-roman.conf").read_text(encoding="utf-8")

        def match(extra: str) -> str:
            with tempfile.TemporaryDirectory() as tmp:
                main = Path(tmp) / "fonts.conf"
                main.write_text('<?xml version="1.0"?><!DOCTYPE fontconfig SYSTEM "urn:fontconfig:fonts.dtd"><fontconfig>'
                                '<dir>/usr/local/share/fonts/bo19</dir><dir>/usr/share/fonts/truetype/dejavu</dir>'
                                f'<cachedir>{tmp}/cache</cachedir>{extra}</fontconfig>', encoding="utf-8")
                out = subprocess.run(["fc-match", "-f", "%{family}", "Times New Roman"], capture_output=True, text=True, timeout=30, env={**os.environ, "FONTCONFIG_FILE": str(main)})
                return out.stdout.strip()

        inner = re.search(r"<fontconfig>(.*)</fontconfig>", conf, re.S).group(1)
        self.assertEqual(match(inner), "Liberation Serif")
        self.assertNotEqual(match(inner.replace("<string>Liberation Serif</string>", "<string>DejaVu Serif</string>")), "Liberation Serif")
        self.assertNotEqual(match(""), "Liberation Serif")  # không có quy tắc nào: hệ điều hành không tự ánh xạ (30-metric-aliases không có trong cấu hình cô lập)


class FileFontVaConf(unittest.TestCase):
    def reference_hashes(self) -> dict[str, str]:
        text = REFERENCE.read_text(encoding="utf-8")
        return dict(re.findall(r"\| `(LiberationSerif-[A-Za-z]+\.ttf)` \| `([0-9a-f]{64})` \|", text))

    def test_bon_file_font_khop_sha256_cua_tai_lieu_tham_chieu(self):
        ref = self.reference_hashes()
        self.assertEqual(set(ref), {"LiberationSerif-Regular.ttf", "LiberationSerif-Bold.ttf", "LiberationSerif-Italic.ttf", "LiberationSerif-BoldItalic.ttf"})
        for name, sha in ref.items():
            self.assertEqual(hashlib.sha256((FONT_DIR / name).read_bytes()).hexdigest(), sha, name)

    def test_co_giay_phep_ofl_di_kem(self):
        text = (FONT_DIR / "LICENSE").read_text(encoding="utf-8")
        self.assertIn("SIL OPEN FONT LICENSE Version 1.1", text)
        ref = re.search(r"\| `LICENSE` \| `([0-9a-f]{64})` \|", REFERENCE.read_text(encoding="utf-8")).group(1)
        self.assertEqual(hashlib.sha256((FONT_DIR / "LICENSE").read_bytes().replace(b"\r\n", b"\n")).hexdigest(), ref)

    def test_file_conf_anh_xa_dung_nhung_gi_font_aliases_noi(self):
        root = etree.parse(str(FONT_DIR / "50-bo19-times-new-roman.conf")).getroot()
        pairs = {}
        for m in root.iter("match"):
            src = m.find("test/string").text
            dst = m.find("edit/string").text
            pairs[src] = dst
            self.assertEqual(m.get("target"), "pattern")
            self.assertEqual(m.find("edit").get("mode"), "assign")
        self.assertEqual(pairs, fonts.FONT_ALIASES)

    def test_dockerfile_copy_dung_file_font_va_conf(self):
        text = (REPO / "Dockerfile").read_text(encoding="utf-8")
        for name in ("LiberationSerif-Regular.ttf", "LiberationSerif-Bold.ttf", "LiberationSerif-Italic.ttf", "LiberationSerif-BoldItalic.ttf", "LICENSE"):
            self.assertIn(f"fonts/{name}", text)
        self.assertIn("fonts/50-bo19-times-new-roman.conf /etc/fonts/conf.d/", text)


class ThuTuStageCuaDockerfile(unittest.TestCase):
    """PO, 2026-10-13: `runtime` LUÔN là stage cuối — Render dựng không `--target`. CI còn kiểm lại bằng cách dựng thật."""

    def stages(self) -> list[tuple[str, str]]:
        text = (REPO / "Dockerfile").read_text(encoding="utf-8")
        return re.findall(r"^FROM\s+(\S+)\s+AS\s+(\S+)\s*$", text, re.M)

    def test_runtime_la_stage_cuoi(self):
        stages = self.stages()
        self.assertEqual([name for _, name in stages], ["base", "test", "runtime"])
        self.assertEqual(stages[-1][1], "runtime")

    def test_test_khong_nam_sau_runtime_va_runtime_khong_dua_tren_test(self):
        stages = dict((name, base) for base, name in self.stages())
        self.assertEqual(stages["runtime"], "base")
        self.assertEqual(stages["test"], "base")

    def test_runtime_khong_cai_poppler_hay_lock_dev(self):
        text = chr(10).join(l for l in (REPO / "Dockerfile").read_text(encoding="utf-8").splitlines() if not l.lstrip().startswith("#"))  # bỏ dòng chú thích: chúng nhắc tên gói
        runtime_part = text[text.index("AS runtime"):]
        self.assertNotIn("poppler", runtime_part)
        self.assertNotIn("requirements-dev", runtime_part)
        base_part = text[: text.index("AS test")]
        self.assertNotIn("poppler", base_part)
        self.assertNotIn("requirements-dev", base_part)

    def test_runtime_chay_bang_user_khong_dac_quyen_va_co_lenh_khoi_dong(self):
        text = (REPO / "Dockerfile").read_text(encoding="utf-8")
        runtime_part = text[text.index("AS runtime"):]
        self.assertIn("USER bo19", runtime_part)
        self.assertIn("CMD", runtime_part)


if __name__ == "__main__":
    unittest.main()

"""Watermark chéo trong PDF thật (B8c; AC-1.3; PO 2026-10-13) — bằng chứng bằng ẢNH TRANG PDF: `soffice` dựng, `pdftoppm` chuyển từng trang sang ảnh xám, và test đọc pixel.

Watermark là shape WordArt nên trong PDF nó là đường vẽ, không phải chữ — `pdftotext` không thấy. Cách kiểm: dựng cùng một văn bản ở `NON_PRODUCTION` và ở `PRODUCTION`, so ảnh từng trang; **những pixel khác nhau CHÍNH LÀ watermark**
(thân văn bản giống hệt nhau vì watermark không đổi bố cục). Điều được chứng minh, nói bằng hành vi:
- **chéo, không phải dòng ngang ở đầu trang:** các pixel khác nhau nằm dọc một đường đi lên từ trái sang phải (hệ số tương quan x–y âm mạnh), trải rộng theo cả hai chiều, có ở hai góc đối nhau của đường chéo và KHÔNG có ở hai góc kia;
- **giữa thân trang:** trọng tâm của các pixel ấy ở giữa trang;
- **chữ thân văn bản vẫn đọc được:** số pixel đậm không giảm, `pdftotext` ra y hệt văn bản không watermark;
- **mọi trang:** văn bản nhiều trang, kể cả có tiêu đề riêng cho trang đầu, nhiều phần, và mẫu không có tiêu đề nào;
- **`PRODUCTION` không có gì:** hai bản dựng ở `PRODUCTION` giống nhau từng pixel và không có đường chéo.

Cần `soffice`, `pdftoppm`, `pdftotext` — image target `test`; thiếu thì bỏ qua, trừ khi `BO19_REQUIRE_SOFFICE=1` (CI).
Chạy từ backend/ với PYTHONPATH=src:  python -m unittest tests.test_watermark_render -v
"""
from __future__ import annotations

from tests import _guard  # noqa: F401 — chốt chặn mạng và khoá API của bộ test (tests/_guard.py)

import io
import statistics
import subprocess
import tempfile
import unittest
from pathlib import Path

from docx import Document

from bo19.tool_layer.rendering import convert, docx_fill
from tests.render_support import DECLARED, VALUES, VIETNAMESE, build_template, pdf_text, require_tool

DPI = 40
DARK = 100  # ngưỡng "chữ đậm" trên thang 0–255


def render(template: bytes, mode: str) -> bytes:
    filled = docx_fill.fill_docx(template, VALUES, declared=DECLARED, operating_mode=mode)
    return convert.convert_to_pdf(filled, timeout_s=300)


def gray_pages(pdf: bytes) -> list[tuple[int, int, bytes]]:
    """Từng trang PDF thành ảnh xám `(rộng, cao, byte)` — PGM nhị phân của `pdftoppm -gray` (không cần thư viện ảnh)."""
    with tempfile.TemporaryDirectory() as d:
        Path(d, "x.pdf").write_bytes(pdf)
        subprocess.run(["pdftoppm", "-gray", "-r", str(DPI), "x.pdf", "pg"], cwd=d, check=True, timeout=120)
        pages = []
        for f in sorted(Path(d).glob("pg-*.pgm")):
            raw = f.read_bytes()
            parts = raw.split(None, 4)  # P5, rộng, cao, 255, rồi dữ liệu — dữ liệu bắt đầu SAU đúng một byte trắng
            w, h = int(parts[1]), int(parts[2])
            pages.append((w, h, raw[-(w * h):]))
        return pages


def changed(a: tuple[int, int, bytes], b: tuple[int, int, bytes], threshold: int = 8) -> list[tuple[int, int]]:
    assert a[:2] == b[:2]
    w = a[0]
    return [(i % w, i // w) for i, (p, q) in enumerate(zip(a[2], b[2])) if abs(p - q) > threshold]


def correlation(points: list[tuple[int, int]]) -> float:
    xs, ys = [p[0] for p in points], [p[1] for p in points]
    return statistics.correlation(xs, ys)


def box(points, w, h, x0, x1, y0, y1) -> int:
    return sum(1 for x, y in points if x0 * w <= x < x1 * w and y0 * h <= y < y1 * h)


def dark(page) -> int:
    return sum(1 for p in page[2] if p < DARK)


def long_template(**kw) -> bytes:
    return build_template(extra_paragraphs=tuple(f"Đoạn {i}: " + VIETNAMESE * 3 for i in range(60)), **kw)


class WatermarkChéo(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        for tool in ("soffice", "pdftoppm", "pdftotext"):
            require_tool(cls, tool)
        cls.template = long_template()
        cls.np_pdf, cls.prod_pdf = render(cls.template, "NON_PRODUCTION"), render(cls.template, "PRODUCTION")
        cls.np_pages, cls.prod_pages = gray_pages(cls.np_pdf), gray_pages(cls.prod_pdf)

    def assert_diagonal(self, points, w, h, label=""):
        self.assertGreater(len(points), 300, f"{label}: quá ít pixel khác nhau — không thấy watermark")
        self.assertLess(correlation(points), -0.7, f"{label}: không phải đường chéo đi lên từ trái sang phải")
        xs, ys = sorted(p[0] for p in points), sorted(p[1] for p in points)
        n = len(points)
        self.assertGreater((xs[int(n * .95)] - xs[int(n * .05)]) / w, 0.5, f"{label}: không trải theo chiều ngang")
        self.assertGreater((ys[int(n * .95)] - ys[int(n * .05)]) / h, 0.3, f"{label}: không trải theo chiều dọc (một dòng ngang)")
        self.assertAlmostEqual(statistics.mean(p[0] for p in points) / w, 0.5, delta=0.12, msg=f"{label}: lệch khỏi giữa trang theo chiều ngang")
        self.assertAlmostEqual(statistics.mean(p[1] for p in points) / h, 0.5, delta=0.12, msg=f"{label}: lệch khỏi giữa trang theo chiều dọc")
        low_left, up_right = box(points, w, h, 0, .3, .62, 1), box(points, w, h, .7, 1, 0, .38)
        up_left, low_right = box(points, w, h, 0, .3, 0, .38), box(points, w, h, .7, 1, .62, 1)
        self.assertGreater(low_left, 30, label)  # hai góc đối nhau của đường chéo có watermark
        self.assertGreater(up_right, 30, label)
        self.assertEqual((up_left, low_right), (0, 0), f"{label}: hai góc kia phải trống")

    def test_van_ban_dai_nhieu_trang(self):
        self.assertGreaterEqual(len(self.np_pages), 3)
        self.assertEqual(len(self.np_pages), len(self.prod_pages))

    def test_mot_duong_cheo_giua_trang_o_trang_dau_va_moi_trang(self):
        for i, (a, b) in enumerate(zip(self.np_pages, self.prod_pages)):
            self.assert_diagonal(changed(a, b), a[0], a[1], f"trang {i + 1}")

    def test_trang_nao_cung_co_ca_trang_cuoi(self):
        counts = [len(changed(a, b)) for a, b in zip(self.np_pages, self.prod_pages)]
        self.assertGreater(min(counts), 0.5 * statistics.median(counts), counts)

    def test_chu_than_van_ban_van_doc_duoc(self):
        for i, (a, b) in enumerate(zip(self.np_pages, self.prod_pages)):
            self.assertGreaterEqual(dark(a), 0.97 * dark(b), f"trang {i + 1}: watermark nuốt chữ đậm")  # watermark nhạt, nằm SAU chữ
        self.assertEqual(" ".join(pdf_text(self.np_pdf).split()), " ".join(pdf_text(self.prod_pdf).split()))  # văn bản trích ra y hệt

    def test_watermark_nhat_hon_chu_nhieu(self):
        a, b = self.np_pages[0], self.prod_pages[0]
        darkest_added = min(p for p, q in zip(a[2], b[2]) if q > 250 and p < 250)  # pixel nền trắng bị watermark phủ: không tối hơn xám trung bình
        self.assertGreater(darkest_added, 120)

    def test_production_khong_co_gi_va_dung_lai_duoc(self):
        again = gray_pages(render(self.template, "PRODUCTION"))
        for a, b in zip(self.prod_pages, again):
            self.assertEqual(changed(a, b, threshold=0), [])  # hai bản dựng PRODUCTION giống hệt từng pixel
        for a, b in zip(self.prod_pages, self.np_pages):
            self.assertEqual(box(changed(a, b), a[0], a[1], 0, 1, 0, .05), 0)  # không có gì ở dải đầu trang: không còn dòng chữ ở tiêu đề như B8b

    def test_tieu_de_rieng_cho_trang_dau(self):
        d = Document(io.BytesIO(self.template))
        d.sections[0].different_first_page_header_footer = True
        d.sections[0].first_page_header.paragraphs[0].text = "Đầu trang đầu"
        buf = io.BytesIO()
        d.save(buf)
        np_pages, prod_pages = gray_pages(render(buf.getvalue(), "NON_PRODUCTION")), gray_pages(render(buf.getvalue(), "PRODUCTION"))
        for i in (0, 1):  # trang đầu dùng tiêu đề trang đầu, trang hai dùng tiêu đề mặc định: cả hai có watermark
            self.assert_diagonal(changed(np_pages[i], prod_pages[i]), np_pages[i][0], np_pages[i][1], f"trang {i + 1}")

    def test_mau_khong_co_tieu_de_nao_van_co_watermark(self):
        template = long_template(header_text=None, footer_text=None)
        np_pages, prod_pages = gray_pages(render(template, "NON_PRODUCTION")), gray_pages(render(template, "PRODUCTION"))
        for i in (0, len(np_pages) - 1):
            self.assert_diagonal(changed(np_pages[i], prod_pages[i], 8), np_pages[i][0], np_pages[i][1], f"trang {i + 1}")

    def test_nhieu_phan_phan_sau_ke_thua_tieu_de(self):
        d = Document(io.BytesIO(build_template(two_sections=True, extra_paragraphs=tuple(VIETNAMESE * 4 for _ in range(30)))))
        buf = io.BytesIO()
        d.save(buf)
        np_pages, prod_pages = gray_pages(render(buf.getvalue(), "NON_PRODUCTION")), gray_pages(render(buf.getvalue(), "PRODUCTION"))
        self.assertGreaterEqual(len(np_pages), 2)
        for i, (a, b) in enumerate(zip(np_pages, prod_pages)):
            self.assert_diagonal(changed(a, b), a[0], a[1], f"trang {i + 1}")


if __name__ == "__main__":
    unittest.main()

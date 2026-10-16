"""Nền chung cho test `document_graph` (B9a, ADR-039): PostgreSQL thật đã migrate, SeaweedFS thật, `soffice` thật, provider LLM là transport giả theo tên lời gọi (`graph_support.FakeLlm`), `DocumentRunner` thật
(checkpointer thật, graph thật, tool thật). Không có test nào ở đây. Mọi nhân viên, tin nhắn, giá trị, mẫu là dữ liệu GIẢ của test.

Mỗi lớp test có DATABASE RIÊNG (template `WORK_CONFIRMATION` đang `ACTIVE` là duy nhất: `template_fetch` đòi đúng một) và mỗi test một bucket.
"""
from __future__ import annotations

import io
import os
import time
import unittest
import uuid
from typing import Any

from docx import Document

from bo19.ai_gateway.budget import Budget
from bo19.ai_gateway.gateway import Gateway
from bo19.ai_gateway.providers import ProviderClient
from bo19.ai_gateway.routing.profiles import load_profiles
from bo19.orchestrator.runner import DocumentRunner, document_id_for
from bo19.persistence.pool import Pool
from bo19.tool_layer import storage as st
from bo19.tool_layer.gate_ops.request_submit import request_submit
from bo19.tool_layer.kernel.permission import employee_context, system_context
from tests import pg_support
from tests.graph_support import KEY, USAGE, FakeLlm  # noqa: F401
from tests.render_support import FONT, require_tool
from tests.request_support import build_request
from tests.s3_support import ENDPOINT, REQUIRED, admin_client, drop_bucket, new_bucket, storage_for

FINGERPRINT_KEY = "KHOA_DAU_VAN_TAY_GIA_CUA_TEST_B9A_DAI_DU_32_BYTE"
NATIONAL_ID = "012345678901"
BODY_OK = "Nhân viên đề nghị cấp giấy xác nhận để bổ sung hồ sơ vay vốn theo yêu cầu của đơn vị tiếp nhận."
SRC_PURPOSE = "TEMPLATE_SOURCE"

# (tên biến, loại, slot nguồn, điền sau duyệt, hướng dẫn, max_length, slot input)
VARIABLES: tuple[tuple[str, str, str | None, bool, str | None, int | None, tuple[str, ...]], ...] = (
    ("full_name", "DIRECT_SLOT", "full_name", False, None, None, ()),
    ("department_name", "DIRECT_SLOT", "department_name", False, None, None, ()),
    ("job_title", "DIRECT_SLOT", "job_title", False, None, None, ()),
    ("recipient_org", "DIRECT_SLOT", "recipient_org", False, None, None, ()),
    ("purpose_statement", "FREE_CONTENT", None, False, "Một câu hành chính ngắn nêu mục đích xin giấy xác nhận.", 400, ("purpose",)),
    ("document_number", "SYSTEM", None, True, None, None, ()),
    ("issued_date", "SYSTEM", None, True, None, None, ()),
    ("signer_user_id", "SYSTEM", None, True, None, None, ()),
)


def template_docx(extra: tuple[str, ...] = ()) -> bytes:
    """`.docx` giả: Times New Roman, có đủ biến của `VARIABLES`, kèm nhãn giả. `extra`: đoạn văn thêm (ví dụ một biến ngoài danh mục)."""
    doc = Document()
    for style in ("Normal", "Header", "Footer"):
        doc.styles[style].font.name = FONT
    doc.add_paragraph("MẪU GIẢ — chỉ để thử, không phải mẫu của tổ chức")
    doc.add_paragraph("{{ full_name }} — {{ department_name }} — {{ job_title }}")
    doc.add_paragraph("Kính gửi: {{ recipient_org }}")
    doc.add_paragraph("{{ purpose_statement }}")
    doc.add_paragraph("Số: {{ document_number }} ngày {{ issued_date }} — người ký {{ signer_user_id }}")
    for text in extra:
        doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def make_template(db, pool, storage, *, status: str = "ACTIVE", variables=VARIABLES, fonts: tuple[str, ...] = (FONT,), data: bytes | None = None) -> uuid.UUID:
    """Template giả cho `WORK_CONFIRMATION` với danh mục biến `variables`. Trả `template_version_id`. Bản gốc ghi qua module lưu trữ (commit thật)."""
    key = f"templates/{uuid.uuid4()}/source.docx"
    st.put_once(pool, storage, object_key=key, purpose=SRC_PURPOSE, content_type="application/octet-stream", data=data or template_docx(), lease_s=120)
    tid, vid = uuid.uuid4(), uuid.uuid4()
    with db.connect(autocommit=True) as c:
        c.execute("set session_replication_role = replica")
        c.execute("insert into template (id, code, request_type_code, name) values (%s, %s, 'WORK_CONFIRMATION', 'MẪU GIẢ')", (tid, f"GIA_{uuid.uuid4().hex[:8]}"))
        c.execute("insert into template_version (id, template_id, version_no, status, source_object_key, uploaded_by_employee_id, required_fonts, activated_at, retired_at) "
                  "values (%s, %s, 1, %s, %s, %s, %s, %s, %s)",
                  (vid, tid, status, key, uuid.uuid4(), list(fonts), None if status == "UPLOADED" else "2026-01-01T00:00:00+00", "2026-02-01T00:00:00+00" if status == "RETIRED" else None))
        for name, kind, source, after, guidance, max_len, inputs in variables:
            c.execute("insert into template_variable (template_version_id, variable_name, kind, source_slot_name, fill_after_approval, variable_guidance, max_length) values (%s, %s, %s, %s, %s, %s, %s)",
                      (vid, name, kind, source, after, guidance, max_len))
            for slot in inputs:
                c.execute("insert into template_variable_input (template_version_id, variable_name, slot_name) values (%s, %s, %s)", (vid, name, slot))
    return vid


def retire_active_templates(db) -> None:
    """Gỡ mọi phiên bản `ACTIVE` đang có trong DB của lớp test (để test dựng template riêng)."""
    with db.connect(autocommit=True) as c:
        c.execute("set session_replication_role = replica")
        c.execute("update template_version set status = 'RETIRED', retired_at = now() where status = 'ACTIVE'")


def draft_output(body: str = BODY_OK, name: str = "purpose_statement") -> dict[str, Any]:
    return {"variable_name": name, "body": body}


@unittest.skipUnless(pg_support.SUPERUSER_DSN, "cần BO19_TEST_PG_SUPERUSER_DSN")
class DocBase(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        if not ENDPOINT:
            if REQUIRED:
                raise AssertionError("BO19_REQUIRE_S3=1 nhưng thiếu BO19_TEST_S3_ENDPOINT")
            raise unittest.SkipTest("cần BO19_TEST_S3_ENDPOINT")
        require_tool(cls, "soffice")
        cls.db = pg_support.PgDb(pg_support.SUPERUSER_DSN)
        cls.pool = Pool(cls.db.dsn("bo19_app"), application_name="bo19-test-doc", min_size=1, max_size=8)
        cls.pool.open()
        cls.llm = FakeLlm()
        client = ProviderClient(base_url="https://llm.example.test/openai/v1", api_key=KEY, transport=cls.llm.transport)
        cls.gateway = Gateway(profiles=load_profiles().profiles, client=client, budget=Budget(cls.pool))

    @classmethod
    def tearDownClass(cls):
        cls.pool.close()
        cls.db.drop()

    def setUp(self):
        self.sys = system_context()
        self.bucket = new_bucket()
        self.addCleanup(drop_bucket, self.bucket)
        self.storage = storage_for(self.bucket)
        self.raw = admin_client()
        self.llm.script.clear()
        self.llm.delays.clear()
        self.llm.requests.clear()
        retire_active_templates(self.db)
        self.version_id = make_template(self.db, self.pool, self.storage)
        self.runner = DocumentRunner.open(pool=self.pool, gateway=self.gateway, storage=self.storage, dsn=self.db.dsn("bo19_app"), fingerprint_key=FINGERPRINT_KEY)
        self.addCleanup(self.runner.close)
        self.eid, _ = self.db.make_employee(roles=("EMPLOYEE",))
        with self.pool.acquire() as conn:
            self.ctx = employee_context(conn, self.eid)

    # ---- dựng dữ liệu -------------------------------------------------------------------------------------------------------

    def submitted_request(self, *, national_id: str | None = None) -> uuid.UUID:
        """Một `request` `WORK_CONFIRMATION` đã đủ điều kiện và đã `SUBMITTED` (qua `request_submit` thật, kèm job `render_document`). `national_id`: thêm slot `RES` `HR_PROFILE` đã xác nhận."""
        rid = build_request(self.pool, self.db, self.eid, self.ctx)
        if national_id is not None:
            with self.db.connect(autocommit=True) as c:
                c.execute("insert into request_slot (request_id, request_type_code, slot_name, value, value_status, provenance_source, provenance_synced_at, confirmed_at) "
                          "values (%s, 'WORK_CONFIRMATION', 'national_id', to_jsonb(%s::text), 'CONFIRMED', 'TEST_FAKE', now(), now())", (rid, national_id))
        with self.db.connect("bo19_migrator") as c:
            version = c.execute("select row_version from request where id = %s", (rid,)).fetchone()[0]
        request_submit(self.pool, self.ctx, request_id=rid, decision_id=uuid.uuid4(), expected_row_version=version)
        return rid

    def run_render(self, rid: uuid.UUID, **kw):
        return self.runner.run_render(request_id=rid, **kw)

    # ---- đọc kiểm (bằng chủ bảng) -------------------------------------------------------------------------------------------

    def scalar(self, sql: str, *args):
        with self.db.connect("bo19_migrator") as c:
            row = c.execute(sql, args).fetchone()
        return row[0] if row else None

    def rows(self, sql: str, *args) -> list[tuple]:
        with self.db.connect("bo19_migrator") as c:
            return c.execute(sql, args).fetchall()

    def document(self, rid: uuid.UUID) -> dict:
        with self.db.connect("bo19_migrator") as c:
            cur = c.execute("select status, revision_round, operating_mode, requires_seal, seal_type, current_draft_render_id, template_version_id from document where id = %s", (document_id_for(rid),))
            row = cur.fetchone()
            return dict(zip([d.name for d in cur.description], row)) if row else {}

    def thread(self, rid: uuid.UUID) -> dict:
        with self.db.connect("bo19_migrator") as c:
            cur = c.execute("select status, waiting_at_node, ended_at from graph_thread where thread_id = %s", (f"document:{document_id_for(rid)}",))
            row = cur.fetchone()
            return dict(zip([d.name for d in cur.description], row)) if row else {}

    def free_content(self, rid: uuid.UUID) -> list[tuple]:
        return self.rows("select variable_name, revision_round, attempt_no, validation_outcome, prompt_module_version, body from document_free_content where document_id = %s order by attempt_no", document_id_for(rid))

    def llm_calls(self) -> int:
        return len(self.llm.calls("draft_free_content"))

    def approve(self, rid: uuid.UUID) -> uuid.UUID:
        """Ghi một `decision_record` `APPROVED` (bằng chủ bảng, thay cho thao tác cổng của B9b) trên bước `CONTENT_REVIEW` đang mở — để thử resume."""
        did = document_id_for(rid)
        approver, _ = self.db.make_employee(roles=("ADMIN_OFFICER",))
        decision = uuid.uuid4()
        with self.db.connect("bo19_migrator") as c:
            step = c.execute("select id from approval_step where document_id = %s and step_kind = 'CONTENT_REVIEW' and status = 'OPEN'", (did,)).fetchone()[0]
            c.execute("insert into decision_record (id, kind, actor_employee_id, request_id, document_id, approval_step_id) values (%s, 'APPROVED', %s, %s, %s, %s)", (decision, approver, rid, did, step))
        return decision

    def checkpoint_text(self) -> str:
        """Mọi dòng của mọi bảng checkpoint của LangGraph, nối thành một chuỗi — để quét giá trị `RES`/văn bản LLM không lọt vào."""
        out = []
        with self.db.connect("bo19_migrator") as c:
            for table in ("checkpoints", "checkpoint_blobs", "checkpoint_writes"):
                for row in c.execute(f"select * from {table}").fetchall():
                    out.append(repr(row))
        return "\n".join(out)

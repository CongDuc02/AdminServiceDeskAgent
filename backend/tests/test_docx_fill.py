"""`bo19.tool_layer.rendering.docx_fill` (B8b) — điền biến, kiểm thiếu/thừa, watermark. Hàm thuần trên byte: không DB, không object storage, không `soffice`.

Điều được chứng minh, nói bằng hành vi:
- **biến thiếu / thừa:** `MISSING_VARIABLE` khi template dùng biến mà bảng không có (thân văn bản, tiêu đề, chân trang); `UNKNOWN_VARIABLE` khi bảng có biến mà phiên bản không khai; lỗi chỉ mang TÊN biến, không giá trị;
- **giá trị là dữ liệu:** `{{`, `{%`, `<`, `&` trong giá trị không bị thông dịch lại hay phá file; ký tự XML cấm bị từ chối; xuống dòng thành ngắt dòng; tiếng Việt đủ dấu giữ nguyên;
- **watermark (AC-1.3):** `NON_PRODUCTION` có dòng watermark ở MỌI tiêu đề (mặc định, trang đầu, trang chẵn, phần sau kế thừa); `PRODUCTION` thì không; chế độ lạ là lỗi lập trình; hàm không có tham số tắt watermark.

Chạy từ backend/ với PYTHONPATH=src:  python -m unittest tests.test_docx_fill -v
"""
from __future__ import annotations

from tests import _guard  # noqa: F401 — chốt chặn mạng và khoá API của bộ test (tests/_guard.py)

import inspect
import io
import unittest

from docx import Document

from bo19.tool_layer.rendering import docx_fill as df
from tests.render_support import DECLARED, FAKE_LABEL, VALUES, VIETNAMESE, build_template, docx_text

NP, PROD = "NON_PRODUCTION", "PRODUCTION"


def fill(template=None, values=None, *, declared=DECLARED, mode=PROD):
    return df.fill_docx(template or build_template(), VALUES if values is None else values, declared=declared, operating_mode=mode)


class DienBien(unittest.TestCase):
    def test_dien_het_bien_o_than_van_ban_tieu_de_va_chan_trang(self):
        text = docx_text(fill())
        for v in VALUES.values():
            self.assertIn(v, text)
        self.assertNotIn("{{", text)
        self.assertIn(FAKE_LABEL, text)  # phần không phải biến giữ nguyên (khung thể thức nằm trong template — ADR-001)

    def test_template_dung_cac_bien_nao(self):
        self.assertEqual(df.template_variables(build_template()), frozenset(DECLARED))
        self.assertEqual(df.template_variables(build_template(header_text=None, footer_text=None)), frozenset({"tieu_de", "ho_ten", "noi_dung"}))

    def test_tieng_viet_du_dau_giu_nguyen(self):
        self.assertIn(VIETNAMESE, docx_text(fill()))

    def test_gia_tri_la_du_lieu_khong_bi_thong_dich_lai(self):
        nasty = "{{ ho_ten }} {% for x in y %} <w:p/> & &amp; ]]> \"'"
        text = docx_text(fill(values={**VALUES, "noi_dung": nasty}))
        self.assertIn(nasty, text)  # nguyên văn: không bị điền lại, không bị escape đôi, không phá XML

    def test_xuong_dong_thanh_ngat_dong(self):
        out = fill(values={**VALUES, "noi_dung": "Dòng một\nDòng hai"})
        self.assertIn("Dòng một", docx_text(out))
        self.assertIn("Dòng hai", docx_text(out))
        self.assertIn(b"<w:br/>", self.document_xml(out))

    @staticmethod
    def document_xml(docx: bytes) -> bytes:
        import zipfile
        return zipfile.ZipFile(io.BytesIO(docx)).read("word/document.xml")

    def test_ket_qua_la_file_word_doc_duoc(self):
        Document(io.BytesIO(fill(mode=NP)))  # không ném


class BienThieuVaThua(unittest.TestCase):
    def test_thieu_bien_o_than_van_ban(self):
        with self.assertRaises(df.MissingVariable) as cm:
            fill(values={k: v for k, v in VALUES.items() if k != "ho_ten"})
        self.assertEqual(cm.exception.names, ("ho_ten",))
        self.assertEqual(cm.exception.code, "MISSING_VARIABLE")

    def test_thieu_bien_o_tieu_de_va_chan_trang(self):
        with self.assertRaises(df.MissingVariable) as cm:
            fill(values={k: v for k, v in VALUES.items() if k not in ("co_quan", "noi_nhan")})
        self.assertEqual(cm.exception.names, ("co_quan", "noi_nhan"))  # docxtpl im lặng điền rỗng — tool thì không

    def test_gia_tri_rong_khong_phai_thieu(self):
        text = docx_text(fill(values={**VALUES, "noi_dung": ""}))
        self.assertNotIn(VIETNAMESE, text)

    def test_bang_co_bien_ma_phien_ban_khong_khai(self):
        with self.assertRaises(df.UnknownVariable) as cm:
            fill(values={**VALUES, "bien_la": "x"})
        self.assertEqual(cm.exception.names, ("bien_la",))
        self.assertEqual(cm.exception.code, "UNKNOWN_VARIABLE")

    def test_bien_khai_ma_template_khong_dung_khong_phai_loi(self):
        fill(declared=(*DECLARED, "khai_nhung_khong_dung"))

    def test_loi_chi_mang_ten_khong_mang_gia_tri(self):
        secret = "GIÁ-TRỊ-BÍ-MẬT-12345"
        for values in ({"ho_ten": secret}, {**VALUES, "bien_la": secret}):
            with self.assertRaises(df.VariableError) as cm:
                fill(values=values)
            blob = f"{cm.exception!r} {cm.exception} {cm.exception.names}"
            self.assertNotIn(secret, blob)

    def test_thieu_duoc_bao_truoc_thua(self):
        with self.assertRaises(df.MissingVariable):
            fill(values={"bien_la": "x"})

    def test_ten_bien_sai_khuon_bi_tu_choi(self):
        with self.assertRaises(df.InvalidVariableValue):
            fill(values={**VALUES, "Ten_Hoa": "x"})

    def test_gia_tri_khong_phai_str_hoac_co_ky_tu_xml_cam_bi_tu_choi(self):
        for bad in (None, 5, b"x", "a\x00b", "a\x0bb", "a\x1fb", "a￾b"):
            with self.assertRaises(df.InvalidVariableValue, msg=repr(bad)) as cm:
                fill(values={**VALUES, "noi_dung": bad})
            self.assertEqual(cm.exception.names, ("noi_dung",))
        fill(values={**VALUES, "noi_dung": "tab\tvà\nxuống dòng thì được"})

    def test_template_hong_la_mot_ma_khong_mang_noi_dung(self):
        for blob in (b"khong phai zip", build_template()[:200], build_template(extra_paragraphs=("{% for x in %}",))):
            with self.assertRaises(df.TemplateUnreadable) as cm:
                fill(template=blob)
            self.assertNotIn("for x", f"{cm.exception!r}")


class Watermark(unittest.TestCase):
    def count(self, docx: bytes) -> int:
        """Số lần chữ watermark xuất hiện trong các PHẦN tiêu đề của file (không đếm hai lần một tiêu đề được hai phần cùng kế thừa)."""
        import zipfile
        zf = zipfile.ZipFile(io.BytesIO(docx))
        return sum(zf.read(n).decode("utf-8").count(df.WATERMARK_TEXT) for n in zf.namelist() if n.startswith("word/header"))

    def test_nguyen_van_ac_1_3(self):
        self.assertEqual(df.WATERMARK_TEXT, "BẢN THỬ NGHIỆM — KHÔNG CÓ GIÁ TRỊ PHÁP LÝ")

    def test_non_production_co_watermark_trong_file(self):
        self.assertEqual(self.count(fill(mode=NP)), 1)

    def test_production_khong_them_gi(self):
        self.assertEqual(self.count(fill(mode=PROD)), 0)

    def test_watermark_o_dau_tieu_de_va_giu_nguyen_noi_dung_tieu_de(self):
        header = Document(io.BytesIO(fill(mode=NP))).sections[0].header
        self.assertEqual([p.text for p in header.paragraphs], [df.WATERMARK_TEXT, "CÔNG TY GIẢ"])

    def test_template_khong_co_tieu_de_van_duoc_watermark(self):
        out = fill(template=build_template(header_text=None, footer_text=None), values={k: v for k, v in VALUES.items() if k in ("tieu_de", "ho_ten", "noi_dung")},
                   declared=("tieu_de", "ho_ten", "noi_dung"), mode=NP)
        self.assertEqual(self.count(out), 1)

    def test_trang_dau_va_trang_chan_deu_co(self):
        template = Document(io.BytesIO(build_template()))
        section = template.sections[0]
        section.different_first_page_header_footer = True
        section.first_page_header.paragraphs[0].text = "Đầu trang đầu"
        template.settings.odd_and_even_pages_header_footer = True
        section.even_page_header.paragraphs[0].text = "Đầu trang chẵn"
        buf = io.BytesIO()
        template.save(buf)
        out = fill(template=buf.getvalue(), mode=NP)
        document = Document(io.BytesIO(out))
        s = document.sections[0]
        for header in (s.header, s.first_page_header, s.even_page_header):
            self.assertEqual(header.paragraphs[0].text, df.WATERMARK_TEXT)

    def test_phan_sau_ke_thua_tieu_de_phan_dau_va_khong_bi_nhan_doi(self):
        out = fill(template=build_template(two_sections=True), mode=NP)
        document = Document(io.BytesIO(out))
        self.assertEqual(len(document.sections), 2)
        self.assertTrue(document.sections[1].header.is_linked_to_previous)
        self.assertEqual(self.count(out), 1)

    def test_phan_sau_co_tieu_de_rieng_cung_co_watermark(self):
        template = Document(io.BytesIO(build_template(two_sections=True)))
        template.sections[1].header.is_linked_to_previous = False
        template.sections[1].header.paragraphs[0].text = "Đầu phần hai"
        buf = io.BytesIO()
        template.save(buf)
        out = fill(template=buf.getvalue(), mode=NP)
        self.assertEqual(self.count(out), 2)

    def test_che_do_la_la_loi_lap_trinh_khong_phai_khong_watermark(self):
        for bad in ("", "production", "TRIAL", None):
            with self.assertRaises(ValueError):
                fill(mode=bad)

    def test_khong_co_tham_so_tat_watermark(self):
        params = set(inspect.signature(df.fill_docx).parameters)
        self.assertEqual(params, {"template", "values", "declared", "operating_mode"})
        for name in params:
            self.assertNotIn("watermark", name.lower())
        self.assertEqual(set(inspect.signature(df.add_watermark).parameters), {"docx"})


if __name__ == "__main__":
    unittest.main()

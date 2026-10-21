"""Dựng mẫu `.docx` GIẤY XÁC NHẬN CÔNG TÁC theo thể thức Phụ lục I Nghị định 30/2020/NĐ-CP, và PDF xem trước bằng dữ liệu GIẢ (PO yêu cầu 2026-10-21).

KHÔNG phải mã ứng dụng, không vào image runtime. Chạy trong image `test` (có python-docx, docxtpl, LibreOffice):

    docker run --rm -v <repo>:/repo -w /repo/backend -e PYTHONPATH=src --entrypoint python bo19-test \
        /repo/tools/template-build/build_work_confirmation.py --out-dir /repo/docs/design/proposals/mau-giay-xac-nhan-cong-tac

Thể thức lấy từ `docs/reference/nghi-dinh-30-2020-phong-chu-va-pham-vi.md` và bản Công báo (Phụ lục I, Phần I mục I–II, Phần V): A4; lề trên/dưới 20–25 mm, trái 30–35 mm, phải 15–20 mm; Times New Roman, màu đen;
Quốc hiệu in hoa đậm 12–13, Tiêu ngữ in thường đậm 13–14 có gạch chân bằng độ dài dòng chữ; tên tổ chức in hoa đậm 12–13 với đường kẻ 1/3–1/2 độ dài; "Số:" cỡ 13; địa danh và ngày nghiêng 13–14;
tên loại văn bản in hoa đậm 13–14; nội dung cỡ 13–14, canh đều, lùi đầu dòng 1 cm hoặc 1,27 cm, dãn dòng từ đơn tới 1,5, giữa đoạn tối thiểu 6 pt; nơi nhận: "Nơi nhận:" cỡ 12 nghiêng đậm, danh sách cỡ 11.
Cỡ chữ thống nhất trong một văn bản (chú thích 1 của bảng Phần V): Quốc hiệu 13, Tiêu ngữ 14, địa danh/ngày 14.

Tên công ty và địa danh là CHỮ CỐ ĐỊNH trong mẫu (không phải biến): PO điền bằng `--company-name`, `--place`. Mặc định là giữ chỗ trong ngoặc vuông — chỉ để xem bố cục.
Biến của mẫu (docxtpl): full_name, department_name, job_title, recipient_org, contract_type (ENUM — chưa có cơ chế trong mã, xem proposals/mau-giay-xac-nhan-cong-tac.md), employment_start_date, employment_end_date,
purpose_statement (FREE_CONTENT), document_number, issued_date, signer_user_id (SYSTEM, điền sau duyệt).
"""
from __future__ import annotations

import argparse
import io
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt

FONT = "Times New Roman"

# (tên biến, giá trị GIẢ) — chỉ dùng cho PDF xem trước; mọi giá trị là dữ liệu giả có nhãn.
FAKE_COMMON = {
    "full_name": "Nguyễn Văn Mẫu (GIẢ)",
    "department_name": "Phòng Mẫu (GIẢ)",
    "job_title": "Chuyên viên (GIẢ)",
    "recipient_org": "Ngân hàng Mẫu (GIẢ)",
    "employment_start_date": "02/03/2022",
    "purpose_statement": "Giấy xác nhận này được cấp để Ông/Bà bổ sung hồ sơ vay vốn tại ngân hàng.",
    "document_number": "01/GXN-ABC",
    "issued_date": "21/10/2026",
    "signer_user_id": "Trần Thị Mẫu (GIẢ), Giám đốc (GIẢ)",
}
# Ba biến thể để duyệt: không xác định thời hạn đang làm; thử việc đang làm; xác định thời hạn đã nghỉ (thể thức "đã từng công tác", EC-WC-02).
SAMPLES = (
    ("INDEFINITE, đang làm việc", {"contract_type": "INDEFINITE", "employment_end_date": ""}),
    ("PROBATION, đang làm việc", {"contract_type": "PROBATION", "employment_end_date": ""}),
    ("FIXED_TERM, đã nghỉ việc", {"contract_type": "FIXED_TERM", "employment_end_date": "15/09/2026"}),
)
VARIABLES = (
    "full_name", "department_name", "job_title", "recipient_org", "contract_type", "employment_start_date", "employment_end_date",
    "purpose_statement", "document_number", "issued_date", "signer_user_id",
)


def _font(run, size: float, *, bold: bool = False, italic: bool = False, underline: bool = False) -> None:
    run.font.name = FONT
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), FONT)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.underline = underline
    run.font.color.rgb = None


def _para(container, text: str = "", size: float = 14, *, bold=False, italic=False, underline=False, align=WD_ALIGN_PARAGRAPH.LEFT, before=0.0, after=0.0, line=1.0, first_indent_cm=0.0):
    p = container.add_paragraph()
    p.alignment = align
    f = p.paragraph_format
    f.space_before, f.space_after, f.line_spacing = Pt(before), Pt(after), line
    if first_indent_cm:
        f.first_line_indent = Cm(first_indent_cm)
    if text:
        _font(p.add_run(text), size, bold=bold, italic=italic, underline=underline)
    return p


def _rule(paragraph, left_cm: float, right_cm: float) -> None:
    """Đường kẻ ngang nét liền dưới đoạn: viền dưới, độ dài do lề đoạn quyết định."""
    paragraph.paragraph_format.left_indent = Cm(left_cm)
    paragraph.paragraph_format.right_indent = Cm(right_cm)
    ppr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "000000")):
        bottom.set(qn(k), v)
    borders.append(bottom)
    ppr.append(borders)


def _no_borders(table) -> None:
    tbl = table._tbl
    tblpr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tblpr.append(borders)


def _cell(table, row: int, col: int, width_cm: float):
    table.columns[col].width = Cm(width_cm)  # LibreOffice đọc độ rộng cột từ `w:tblGrid`, không từ `w:tcW`
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    if table._tbl.tblPr.find(qn("w:tblLayout")) is None:
        table._tbl.tblPr.append(layout)
    c = table.cell(row, col)
    c.width = Cm(width_cm)
    c.paragraphs[0]._element.getparent().remove(c.paragraphs[0]._element)  # ô không đoạn thừa
    return c


def _jinja(p, size: float, *parts: tuple[str, dict]) -> None:
    for text, style in parts:
        _font(p.add_run(text), size, **style)


def build(company_name: str, place: str) -> bytes:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)  # A4
    sec.top_margin, sec.bottom_margin, sec.left_margin, sec.right_margin = Mm(20), Mm(20), Mm(30), Mm(15)
    for name in ("Normal", "Header", "Footer"):  # mọi họ font khai trong file chỉ là Times New Roman (manifest `required_fonts`)
        st = doc.styles[name]
        st.font.name = FONT
        rpr = st.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.insert(0, rfonts)
        for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rfonts.set(qn(attr), FONT)
    body = doc.element.body
    for p in list(doc.paragraphs):  # python-docx mẫu có một đoạn rỗng
        body.remove(p._element)

    text_width = 165  # mm: 210 - 30 - 15
    # --- Đầu trang: tên tổ chức (trái) | Quốc hiệu, Tiêu ngữ (phải) -------------------------------------------------------------
    head = doc.add_table(rows=1, cols=2)
    head.alignment = WD_TABLE_ALIGNMENT.CENTER
    head.autofit = False
    _no_borders(head)
    left, right = _cell(head, 0, 0, 6.0), _cell(head, 0, 1, 10.5)
    org = _para(left, company_name.upper(), 13, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)  # ô 2 — tên tổ chức ban hành: in hoa, đậm, 12–13
    approx_text_cm = max(2.0, len(company_name) * 0.30)  # ~chiều dài dòng chữ ở cỡ 13 in hoa đậm
    line_cm = approx_text_cm * 0.42  # 1/3–1/2 độ dài dòng chữ
    margin = max(0.0, (6.0 - line_cm) / 2)
    rule = _para(left, "", 13, align=WD_ALIGN_PARAGRAPH.CENTER)
    _rule(rule, margin, margin)
    rule.paragraph_format.space_before = Pt(0)
    number = _para(left, "", 13, align=WD_ALIGN_PARAGRAPH.CENTER, before=6)  # ô 3 — Số: cỡ 13, đứng
    _jinja(number, 13, ("Số: ", {}), ("{{ document_number }}", {}))
    _para(right, "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM", 13, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)  # ô 1
    _para(right, "Độc lập - Tự do - Hạnh phúc", 14, bold=True, underline=True, align=WD_ALIGN_PARAGRAPH.CENTER, before=0)
    date = _para(right, "", 14, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, before=6)  # ô 4 — địa danh và ngày: nghiêng, 14
    _jinja(date, 14,
           (f"{place}, ", {"italic": True}),
           ("{% if issued_date[2:3] == '/' %}ngày {{ issued_date[0:2] }} tháng {{ issued_date[3:5] }} năm {{ issued_date[6:10] }}{% else %}{{ issued_date }}{% endif %}", {"italic": True}))

    # --- Tên loại văn bản -----------------------------------------------------------------------------------------------------
    _para(doc, "GIẤY XÁC NHẬN CÔNG TÁC", 14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, before=22, after=14)  # ô 5a
    kg = _para(doc, "", 14, align=WD_ALIGN_PARAGRAPH.LEFT, after=6)  # ô 9a — Kính gửi
    _jinja(kg, 14, ("Kính gửi: ", {}), ("{{ recipient_org }}", {}))

    # --- Nội dung (ô 6): canh đều, lùi đầu dòng 1,27 cm, dãn dòng 1,3, giữa đoạn 6 pt --------------------------------------------
    def line(text_parts, *, indent=1.27, after=6.0):
        p = _para(doc, "", 14, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=after, line=1.3, first_indent_cm=indent)
        _jinja(p, 14, *text_parts)
        return p

    line([("Công ty xác nhận:", {})])
    for label, var in (("Họ và tên", "full_name"), ("Bộ phận công tác", "department_name"), ("Chức danh", "job_title")):
        line([(f"{label}: ", {}), ("{{ " + var + " }}", {"bold": True})], indent=1.27, after=3)
    line([
        ("{% if employment_end_date %}Đã làm việc{% else %}Hiện đang làm việc{% endif %} tại Công ty từ ngày {{ employment_start_date }}"
         "{% if employment_end_date %} đến ngày {{ employment_end_date }}{% endif %}"
         "{% if contract_type == 'PROBATION' %}, trong thời gian thử việc"
         "{% elif contract_type == 'FIXED_TERM' %}, theo hợp đồng lao động xác định thời hạn"
         "{% elif contract_type == 'INDEFINITE' %}, theo hợp đồng lao động không xác định thời hạn{% endif %}.", {}),
    ])
    line([("{{ purpose_statement }}", {})])

    # --- Nơi nhận (ô 9b, trái) | chữ ký (ô 7, phải) ----------------------------------------------------------------------------
    foot = doc.add_table(rows=1, cols=2)
    foot.alignment = WD_TABLE_ALIGNMENT.CENTER
    foot.autofit = False
    _no_borders(foot)
    fl, fr = _cell(foot, 0, 0, 7.5), _cell(foot, 0, 1, 9.0)
    _para(fl, "Nơi nhận:", 12, bold=True, italic=True, before=14)
    _para(fl, "- Như trên;", 11)
    _para(fl, "- Lưu: VT.", 11)
    _para(fr, "(Ký, ghi rõ họ tên, chức danh và đóng dấu)", 13, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, before=14)
    for _ in range(4):  # chỗ ký và đóng dấu
        _para(fr, "", 14)
    signer = _para(fr, "", 14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    _jinja(signer, 14, ("{{ signer_user_id }}", {"bold": True}))
    doc.core_properties.title = "Giấy xác nhận công tác"  # `title` giữ nguyên qua LibreOffice sang PDF (ADR-038, M1.1)
    doc.core_properties.author = "BO-19"
    doc.core_properties.comments = "Mẫu nháp — chưa duyệt"
    out = io.BytesIO()
    doc.save(out)
    return _only_times_new_roman(out.getvalue())


_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_A = "http://schemas.openxmlformats.org/drawingml/2006/main"


def _only_times_new_roman(blob: bytes) -> bytes:
    """Mẫu python-docx mang sẵn bảng font, theme và kiểu có Calibri/Cambria/Arial/…: dọn để họ font khai trong file CHỈ là Times New Roman (manifest `required_fonts`, ADR-015/ADR-036)."""
    from lxml import etree

    src, out = zipfile.ZipFile(io.BytesIO(blob)), io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename in ("word/styles.xml", "word/stylesWithEffects.xml", "word/numbering.xml"):  # numbering.xml của mẫu mang Symbol/Courier cho dấu đầu dòng — mẫu này không dùng danh sách
                root = etree.fromstring(data)
                for rf in root.iter(f"{{{_W}}}rFonts"):
                    for attr in list(rf.attrib):
                        del rf.attrib[attr]
                    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
                        rf.set(f"{{{_W}}}{attr}", FONT)
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            elif item.filename == "word/fontTable.xml":
                data = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:fonts xmlns:w="{_W}"><w:font w:name="{FONT}"><w:family w:val="roman"/></w:font></w:fonts>').encode()
            elif item.filename.startswith("word/theme/") and item.filename.endswith(".xml"):
                root = etree.fromstring(data)
                for kind in ("majorFont", "minorFont"):
                    for node in root.iter(f"{{{_A}}}{kind}"):
                        for child in list(node):
                            if child.tag == f"{{{_A}}}font":
                                node.remove(child)
                        for tag, face in (("latin", FONT), ("ea", ""), ("cs", "")):
                            el = node.find(f"{{{_A}}}{tag}")
                            if el is not None:
                                el.set("typeface", face)
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            dst.writestr(item, data)
    return out.getvalue()


def preview(template: bytes, out_pdf: Path) -> None:
    """PDF xem trước, đi đúng đường của `docx_render`: `fill_docx` (kiểm biến, watermark NON_PRODUCTION) rồi LibreOffice."""
    from bo19.tool_layer.rendering.convert import convert_to_pdf
    from bo19.tool_layer.rendering.docx_fill import fill_docx

    pages: list[bytes] = []
    for _, extra in SAMPLES:
        values = {**FAKE_COMMON, **extra}
        assert set(values) == set(VARIABLES), set(values) ^ set(VARIABLES)
        pages.append(convert_to_pdf(fill_docx(template, values, declared=VARIABLES, operating_mode="NON_PRODUCTION"), timeout_s=120))
    with tempfile.TemporaryDirectory() as tmp:
        files = []
        for i, blob in enumerate(pages):
            f = Path(tmp) / f"p{i}.pdf"
            f.write_bytes(blob)
            files.append(str(f))
        if shutil.which("pdfunite") is None:
            sys.exit("thiếu pdfunite (poppler-utils)")
        subprocess.run(["pdfunite", *files, str(out_pdf)], check=True)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--out-dir", required=True, type=Path)
    ap.add_argument("--company-name", default="[TÊN CÔNG TY]")
    ap.add_argument("--place", default="[Địa danh]")
    ap.add_argument("--no-preview", action="store_true")
    args = ap.parse_args()
    from bo19.tool_layer.rendering.docx_fill import template_variables
    from bo19.tool_layer.rendering.fonts import declared_families

    blob = build(args.company_name, args.place)
    fams = declared_families(blob)
    if fams != {FONT}:
        sys.exit(f"họ font ngoài Times New Roman: {sorted(fams)}")
    used = template_variables(blob)
    if used != set(VARIABLES):
        sys.exit(f"biến của mẫu lệch danh mục: {sorted(used ^ set(VARIABLES))}")
    args.out_dir.mkdir(parents=True, exist_ok=True)
    (args.out_dir / "tpl_work_confirmation.docx").write_bytes(blob)
    print("font:", sorted(fams), "| biến:", sorted(used))
    if not args.no_preview:
        preview(blob, args.out_dir / "xem-truoc-du-lieu-gia.pdf")
        print("đã xuất xem-truoc-du-lieu-gia.pdf (3 trang, dữ liệu giả, watermark NON_PRODUCTION)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

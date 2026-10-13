"""So ngắt dòng và ngắt trang của PDF do `soffice` dựng khi `.docx` khai "Times New Roman" — với font THẬT Times New Roman 7.12 và với Liberation Serif 2.1.5 (B8b; món nợ đo của ADR-036: "kerning và render thật").

KHÔNG phải mã ứng dụng, không vào image. Chỉ chạy trên máy người triển khai: Times New Roman 7.12 là bản cài sẵn của Windows (`C:\\Windows\\Fonts\\times*.ttf`), **chỉ đọc để đo**, gắn chỉ đọc vào container dùng một lần, không chép vào repo hay image.

Cách dùng (từ thư mục gốc repo; image target `test`; `<fonts-win>` chứa times.ttf, timesbd.ttf, timesi.ttf, timesbi.ttf):
    docker run --rm -v "$PWD":/repo -v <fonts-win>:/tnr:ro -w /repo/backend -e PYTHONPATH=src:. --entrypoint python bo19-test /repo/tools/font-probe/compare_linebreaks.py

Hai lần dựng cùng MỘT `.docx` bằng hai cấu hình fontconfig cô lập: (A) chỉ có Times New Roman thật (và DejaVu để dự phòng) — `fc-match "Times New Roman"` ra chính nó; (B) chỉ có Liberation Serif của dự án cộng quy tắc
`fonts/50-bo19-times-new-roman.conf`. Rồi so, bằng `pdftotext -layout`: số trang, số dòng mỗi trang, và **từng dòng chữ**. Hai PDF nhúng hai font khác nhau nên chỉ so CHỮ và CHỖ NGẮT, không so byte.
Văn bản: tiếng Việt đủ dấu ghép từ một kho câu GIẢ bằng bộ sinh số giả ngẫu nhiên cố định hạt giống — nhiều đoạn, bốn kiểu chữ (thường, đậm, nghiêng, đậm-nghiêng), canh đều hai bên và canh trái, hai cỡ chữ.
"""
from __future__ import annotations

import io
import json
import os
import random
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from bo19.tool_layer.rendering import convert

SENTENCES = [
    "Nguyễn Thị Hương Giang đã công tác tại Phòng Hành chính — Quản trị từ ngày 01/01/2020 và được giao quyền ký thay theo văn bản số 123/QĐ-GIẢ.",
    "Đề nghị các đơn vị liên quan phối hợp thực hiện đúng thời hạn, báo cáo kết quả bằng văn bản về Văn phòng trước ngày 15 hằng tháng.",
    "Giấy giới thiệu có giá trị trong thời hạn ba mươi ngày kể từ ngày ký; quá thời hạn trên phải làm lại thủ tục theo quy định hiện hành.",
    "Phòng họp tầng ba được đặt từ 14 giờ 30 đến 16 giờ 00 ngày thứ Năm, sức chứa tối đa hai mươi người, có máy chiếu và bảng trắng.",
    "Chúng tôi xin trân trọng cảm ơn sự hợp tác của Quý cơ quan; mọi chi tiết xin liên hệ số điện thoại nội bộ 0123 theo giờ hành chính.",
    "Việc sử dụng con dấu phải được cấp có thẩm quyền phê duyệt, ghi sổ theo dõi, và chỉ đóng vào văn bản đã có chữ ký hợp lệ.",
    "Những trường hợp đặc biệt như ủy quyền, thay thế, thu hồi văn bản được xử lý riêng, kèm lý do bằng văn bản và lưu hồ sơ gốc.",
]
STYLES = {"thường": (False, False), "đậm": (True, False), "nghiêng": (False, True), "đậm-nghiêng": (True, True)}


def build(size_pt: int, justify: bool, seed: int) -> bytes:
    rng = random.Random(seed)
    d = Document()
    style = d.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(size_pt)
    rfonts = style.element.get_or_add_rPr().find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rFonts")
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(f"{{http://schemas.openxmlformats.org/wordprocessingml/2006/main}}{attr}", "Times New Roman")
    for name, (bold, italic) in STYLES.items():
        d.add_paragraph(f"Kiểu {name}").runs[0].bold = True
        for _ in range(8):
            p = d.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if justify else WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(" ".join(rng.sample(SENTENCES, rng.randint(2, 5))))
            run.bold, run.italic = bold, italic
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def conf(dirs: list[str], extra: str, tmp: str, tag: str) -> str:
    path = Path(tmp) / f"{tag}.conf"
    body = "".join(f"<dir>{x}</dir>" for x in dirs)
    path.write_text('<?xml version="1.0"?><!DOCTYPE fontconfig SYSTEM "urn:fontconfig:fonts.dtd"><fontconfig>' + body + f"<cachedir>{tmp}/cache-{tag}</cachedir>{extra}</fontconfig>", encoding="utf-8")
    return str(path)


def pages(pdf: bytes) -> list[list[str]]:
    out = subprocess.run(["pdftotext", "-layout", "-", "-"], input=pdf, capture_output=True, check=True, timeout=120).stdout.decode("utf-8")
    return [[" ".join(line.split()) for line in page.splitlines() if line.strip()] for page in out.split("\f") if page.strip()]


def fonts_in(pdf: bytes) -> list[str]:
    out = subprocess.run(["pdffonts", "-"], input=pdf, capture_output=True, check=True, timeout=60).stdout.decode("utf-8").splitlines()[2:]
    return sorted({line.split()[0].split("+")[-1] for line in out if line.strip()})


def render(docx: bytes, config: str) -> bytes:
    old = os.environ.get("FONTCONFIG_FILE")
    os.environ["FONTCONFIG_FILE"] = config
    try:
        return convert.convert_to_pdf(docx, timeout_s=600)
    finally:
        if old is None:
            os.environ.pop("FONTCONFIG_FILE")
        else:
            os.environ["FONTCONFIG_FILE"] = old


def main() -> int:
    if not Path("/tnr/times.ttf").exists():
        print("thiếu /tnr/times.ttf — gắn thư mục chứa Times New Roman 7.12 của Windows (chỉ đọc)", file=sys.stderr)
        return 2
    repo_conf = Path("/repo/fonts/50-bo19-times-new-roman.conf").read_text(encoding="utf-8")
    inner = repo_conf[repo_conf.index("<fontconfig>") + len("<fontconfig>"):repo_conf.index("</fontconfig>")]
    results = []
    with tempfile.TemporaryDirectory() as tmp:
        real = conf(["/tnr", "/usr/share/fonts/truetype/dejavu"], "", tmp, "real")
        liberation = conf(["/usr/local/share/fonts/bo19", "/usr/share/fonts/truetype/dejavu"], inner, tmp, "lib")
        for size in (12, 13):
            for justify in (True, False):
                docx = build(size, justify, seed=size * 10 + int(justify))
                a, b = render(docx, real), render(docx, liberation)
                pa, pb = pages(a), pages(b)
                lines_a, lines_b = [l for p in pa for l in p], [l for p in pb for l in p]
                differing = [i for i, (x, y) in enumerate(zip(lines_a, lines_b)) if x != y]
                results.append({"size_pt": size, "justify": justify, "pages": [len(pa), len(pb)], "lines_per_page_equal": [len(p) for p in pa] == [len(p) for p in pb],
                                "lines": [len(lines_a), len(lines_b)], "lines_differing": len(differing) + abs(len(lines_a) - len(lines_b)), "fonts_real": fonts_in(a), "fonts_liberation": fonts_in(b),
                                "first_difference": None if not differing else {"real": lines_a[differing[0]], "liberation": lines_b[differing[0]]}})
    print(json.dumps(results, ensure_ascii=False, indent=1))
    return 0


if __name__ == "__main__":
    sys.exit(main())
